import io
import csv
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from .. import models, auth, analytics
from ..profile_defaults import FIELD_LABELS

router = APIRouter(prefix="/api/studies", tags=["results"])

# Campos por los que se puede segmentar (atributo del modelo Respondent).
SEG_FIELDS = ["sex", "age_group", "occupation", "education",
              "municipality", "district", "electoral_section",
              "locality_zone", "operator"]


def _owned(study_id: str, user, db) -> models.Study:
    study = db.get(models.Study, study_id)
    if not study or study.owner_id != user.id:
        raise HTTPException(404, "Estudio no encontrado")
    return study


def _all_respondents(study, db):
    return db.query(models.Respondent).filter(
        models.Respondent.study_id == study.id).all()


def _segments(respondents):
    """Valores distintos por campo (para los menús de filtro), con conteo."""
    out = []
    for f in SEG_FIELDS:
        counts = {}
        for r in respondents:
            if r.completed_at is None:
                continue
            v = (getattr(r, f, "") or "").strip()
            if v:
                counts[v] = counts.get(v, 0) + 1
        if counts:
            out.append({
                "field": f, "label": FIELD_LABELS.get(f, f),
                "values": [{"value": k, "n": v} for k, v in sorted(counts.items(), key=lambda x: -x[1])],
            })
    return out


def _results(study, db, seg_field=None, seg_value=None):
    respondents = _all_respondents(study, db)
    segments = _segments(respondents)
    subset = respondents
    applied = None
    if seg_field and seg_value and seg_field in SEG_FIELDS:
        subset = [r for r in respondents if (getattr(r, seg_field, "") or "").strip() == seg_value]
        applied = {"field": seg_field, "label": FIELD_LABELS.get(seg_field, seg_field), "value": seg_value}
    res = analytics.compute_results(study, subset)
    res["segments"] = segments
    res["applied_filter"] = applied
    return res


@router.get("/{study_id}/results")
def results(study_id: str, seg_field: str = Query(None), seg_value: str = Query(None),
            user: models.User = Depends(auth.get_current_user), db: Session = Depends(get_db)):
    return _results(_owned(study_id, user, db), db, seg_field, seg_value)


@router.get("/{study_id}/export.csv")
def export_csv(study_id: str, user: models.User = Depends(auth.get_current_user), db: Session = Depends(get_db)):
    study = _owned(study_id, user, db)
    R = _results(study, db)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Atributo", "Categoria", "Veces mostrada", "Veces elegida",
                "% eleccion (del atributo)", "Tasa preferencia %"])
    for a in R["by_attribute"]:
        for c in a["categories"]:
            w.writerow([a["attribute_name"], c["category_name"], c["shown"], c["chosen"],
                        round(c["share"] * 100, 2), round(c["win_rate"] * 100, 2)])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{study.name}_categorias.csv"'},
    )


@router.get("/{study_id}/export.xlsx")
def export_xlsx(study_id: str, user: models.User = Depends(auth.get_current_user), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    study = _owned(study_id, user, db)
    R = _results(study, db)
    respondents = _all_respondents(study, db)
    attrs = study.attributes
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Métrica", "Valor"])
    ws.append(["Estudio", study.name])
    ws.append(["Encuestados con respuesta", R["n_responses"]])
    ws.append(["Total de elecciones", R["total_choices"]])
    ws.append(["Atributos", R["attribute_count"]])
    ws.append(["Exportado", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])

    ws2 = wb.create_sheet("Por categoría")
    ws2.append(["Atributo", "Categoría", "Veces mostrada", "Veces elegida", "% elección", "Tasa preferencia %"])
    for a in R["by_attribute"]:
        for c in a["categories"]:
            ws2.append([a["attribute_name"], c["category_name"], c["shown"], c["chosen"],
                        round(c["share"] * 100, 2), round(c["win_rate"] * 100, 2)])

    ws3 = wb.create_sheet("Ranking categorías")
    ws3.append(["#", "Atributo", "Categoría", "Tasa preferencia %", "Veces elegida", "Veces mostrada"])
    for i, c in enumerate(R["all_categories"], 1):
        ws3.append([i, c["attribute_name"], c["category_name"],
                    round(c["win_rate"] * 100, 2), c["chosen"], c["shown"]])

    ws4 = wb.create_sheet("Combinaciones")
    ws4.append(["#", "Combinación", "Veces mostrada", "Veces elegida", "% del total", "Tasa de victoria %"])
    for i, c in enumerate(R["combos"], 1):
        ws4.append([i, c["label"], c["shown"], c["chosen"],
                    round(c["share"] * 100, 2), round(c["win_rate"] * 100, 2)])

    # Hoja: entrevistados (perfil completo, una fila por persona)
    wsR = wb.create_sheet("Entrevistados")
    wsR.append(["ID entrevistado", "Fecha y hora", "Encuestador", "Sexo", "Edad",
                "Ocupación", "Escolaridad", "Municipio", "Distrito",
                "Sección electoral", "Localidad / colonia / zona"])
    for r in respondents:
        wsR.append([
            r.id, r.completed_at.strftime("%Y-%m-%d %H:%M") if r.completed_at else "",
            r.operator or "", r.sex or "", r.age_group or "", r.occupation or "",
            r.education or "", r.municipality or "", r.district or "",
            r.electoral_section or "", r.locality_zone or "",
        ])

    # Hoja: una fila por interacción con perfil + alternativa elegida
    ws5 = wb.create_sheet("Por interacción")
    ws5.append(["ID entrevistado", "Encuestador", "Sexo", "Edad", "Ocupación", "Escolaridad",
                "Municipio", "Distrito", "Sección electoral", "Localidad",
                "Interacción"] + [a.name for a in attrs])
    for r in respondents:
        base = [r.id, r.operator or "", r.sex or "", r.age_group or "", r.occupation or "",
                r.education or "", r.municipality or "", r.district or "",
                r.electoral_section or "", r.locality_zone or ""]
        for inter in sorted(r.interactions, key=lambda x: x.task_index):
            chosen = next((o for o in inter.options
                           if o.option_index == inter.chosen_option_index), None)
            by_attr = {it.attribute_id: it.category_name for it in chosen.items} if chosen else {}
            ws5.append(base + [inter.task_index + 1] + [by_attr.get(a.id, "") for a in attrs])

    # Hoja: resumen por encuestador
    ws6 = wb.create_sheet("Por operador")
    ws6.append(["Operador", "Entrevistas completadas", "Elecciones registradas"])
    for o in R.get("by_operator", []):
        ws6.append([o["operator"], o["respondents"], o["choices"]])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return StreamingResponse(
        bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{study.name}_resultados.xlsx"'},
    )
